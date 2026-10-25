"""Generate synthetic upload examples only; no customer files are read."""
import datetime
import pathlib
import sys
from openpyxl import Workbook
from docx import Document
from reportlab.pdfgen import canvas

destination = pathlib.Path(sys.argv[1])
destination.mkdir(parents=True, exist_ok=True)
(destination / "报价记录.csv").write_text('地点,品种,价格,备注\n演示地点甲,演示商品A,685,"含税,演示单位"\n演示地点乙,演示商品B,580,"第一行\n第二行"\n', encoding="utf-8-sig")
(destination / "中文编码.csv").write_bytes("品种,报价\n演示商品C,3500\n".encode("gb18030"))
(destination / "说明.md").write_text("# 评测参考资料\n\n本资料用于检验文件预览。\n\n## 评分要求\n\n- 数值准确\n- 口径一致\n\n| 品种 | 价格 |\n| --- | --- |\n| 演示商品A | 685 |\n\n<script>window.injected=true</script>\n", encoding="utf-8")
(destination / "备注.txt").write_text("模型评测说明\n第一条：按提供的数据回答。\n第二条：保留日期和单位。", encoding="utf-8")
(destination / "UTF16.txt").write_text("中文 UTF16 文件", encoding="utf-16")
(destination / "空文件.txt").write_bytes(b"")
(destination / "附件.bin").write_bytes(bytes([0, 1, 2, 3, 255]))
(destination / "损坏.docx").write_bytes(b"not a zip archive")
(destination / "旧版.doc").write_bytes(b"legacy document fixture")
(destination / "数据.tsv").write_text("编号\t答案\n001\t正确\n", encoding="utf-8")
book = Workbook()
sheet = book.active
sheet.title = "演示数据"
sheet.append(["日期", "地点", "品种", "价格"])
sheet.append([datetime.datetime(2026, 8, 19), "演示地点甲", "演示商品A", 685])
second = book.create_sheet("专家意见")
second.append(["题目ID", "专家结论"])
second.append(["IO-001", "回答应明确日期、规格、单位与含税口径"])
book.save(destination / "评测样例.xlsx")
large = Workbook()
for index in range(250):
    large.active.append([index, "测试数据"])
large.save(destination / "多行.xlsx")
doc = Document()
doc.add_heading("虚构问题与专家意见", level=1)
doc.add_paragraph("本文件包含需要用于评测的问题和专家意见。")
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "问题"
table.cell(0, 1).text = "专家意见"
table.cell(1, 0).text = "演示地点甲演示商品A价格是多少？"
table.cell(1, 1).text = "需包含规格与报价口径。"
doc.add_paragraph("表格后的说明。")
doc.save(destination / "专家意见.docx")
pdf = canvas.Canvas(str(destination / "附件.pdf"))
pdf.drawString(72, 750, "PDF preview test")
pdf.save()
print(str(destination))
