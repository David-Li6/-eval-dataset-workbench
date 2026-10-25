import csv
import json
import pathlib
import sys
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

root = pathlib.Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
workbook = Workbook()
sheet = workbook.active
sheet.title = "原题"
sheet.append(["题目id", "问题", "专家结论/意见"])
for number in range(1, 252):
    sheet.append([f"IO-{number:03}", f"问题{number}", "专家长意见" * 700 if number == 251 else "必须回答单位"])
other = workbook.create_sheet("补充意见")
other.append(["编号", "口径落点提示"])
other.append(["EX-001", "第二工作表末尾标记"])
workbook.save(root / "完整表格.xlsx")
with open(root / "多行编码.csv", "w", encoding="gb18030", newline="") as stream:
    writer = csv.writer(stream)
    writer.writerow(["id", "专家意见"])
    writer.writerow(["CSV-001", "包含逗号,以及\n换行中文"])
document = Document()
document.add_paragraph("前文标记" + "完整正文" * 4000)
table = document.add_table(rows=2, cols=2)
table.cell(0, 0).text, table.cell(0, 1).text = "编号", "专家意见"
table.cell(1, 0).text, table.cell(1, 1).text = "DOC-001", "表格结论标记"
document.add_paragraph("末尾正文标记")
document.save(root / "完整文档.docx")
pdf = canvas.Canvas(str(root / "全部页面.pdf"))
for page in range(1, 26):
    pdf.drawString(50, 750, f"Source page {page} complete marker")
    pdf.showPage()
pdf.save()
(root / "长文本.txt").write_text("原始正文\n" * 45000 + "长文本末尾标记", encoding="utf-8")
(root / "记录.json").write_text(json.dumps({"questions": [{"题目id": "JSON-001", "专家意见": "完整 JSON 记录"}]}, ensure_ascii=False), encoding="utf-8")
(root / "记录.jsonl").write_text('\n'.join(json.dumps({"id": f"LINE-{n}", "意见": f"第{n}条"}, ensure_ascii=False) for n in range(1, 4)), encoding="utf-8")
(root / "空文件.txt").write_text("", encoding="utf-8")
(root / "超限.txt").write_text("字" * 2000001, encoding="utf-8")
