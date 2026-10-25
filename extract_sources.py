"""Read complete source text within explicit limits; never execute document contents."""
import csv
import datetime
import hashlib
import io
import json
import pathlib
import re
import sys
import unicodedata
from preview_file import zip_guard, PreviewError, TEXT_EXT

LIMIT = 2000000
MAX_UNIT = 60000
units, total = [], 0


def cell(value):
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)


def add(text, location, original_id=""):
    global total
    if not text.strip():
        return
    if len(text) > MAX_UNIT:
        raise PreviewError(f"{location} 内容超过 60,000 字符，请拆分该行或段落后重试。")
    total += len(text)
    if total > LIMIT or len(units) >= 10000:
        raise PreviewError("源文档内容超过生成范围（每文件 200 万字符、10,000 段），请拆分文件。未截断正文。")
    units.append({"text": text, "location": location, "originalId": original_id})


def normalized(value):
    return re.sub(r"[\s_-]", "", unicodedata.normalize("NFKC", value).lower())


def table(rows, prefix):
    header, header_number, id_column, data_count = None, None, None, 0
    for number, row in enumerate(rows, 1):
        values = [cell(value) for value in row]
        while values and not values[-1]:
            values.pop()
        if not any(value.strip() for value in values):
            continue
        if header is None:
            header, header_number = values, number
            id_column = next((i for i, value in enumerate(values) if normalized(value) in {"id", "题目id", "问题id", "编号", "题号", "来源id"}), None)
            continue
        lines = [f"{header[i] if i < len(header) and header[i].strip() else '列' + str(i + 1)}：{value}" for i, value in enumerate(values) if value]
        add("\n".join(lines), f"{prefix}第 {number} 行", values[id_column] if id_column is not None and id_column < len(values) else "")
        data_count += 1
    if header and not data_count:
        add("\n".join(header), f"{prefix}第 {header_number} 行")


def decode(filename):
    raw = pathlib.Path(filename).read_bytes()
    encoding = "utf-32" if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")) else "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        text = raw.decode("gb18030")
    if "\x00" in text or len(text) > LIMIT:
        raise PreviewError("文档不是可读取文本或超过 200 万字符，请检查编码或拆分后重试。")
    return text


def paragraphs(text, prefix):
    # Split long plain-text paragraphs, retaining every character and order.
    start = 0
    while start < len(text):
        end = min(start + 12000, len(text))
        if end < len(text):
            boundary = text.rfind("\n", start + 6000, end)
            if boundary >= 0:
                end = boundary + 1
        add(text[start:end], f"{prefix}字符 {start + 1}–{end}")
        start = end


def main():
    filename, name = sys.argv[1:3]
    extension = pathlib.Path(name).suffix.lower()
    notes = []
    if extension in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        zip_guard(filename)
        with open(filename, "rb") as stream:
            workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in workbook:
                    if sheet.max_row == 1 and sheet.max_column == 1:
                        sheet.reset_dimensions()
                    table(sheet.iter_rows(values_only=True), f"工作表「{sheet.title}」")
            finally:
                workbook.close()
        notes.append("已读取全部工作表及长单元格；公式保留原文，不执行宏。")
    elif extension in {".csv", ".tsv"}:
        text = decode(filename)
        try:
            dialect = csv.Sniffer().sniff(text[:16000], delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel_tab if extension == ".tsv" else csv.excel
        csv.field_size_limit(LIMIT)
        table(csv.reader(io.StringIO(text, newline=""), dialect), "")
    elif extension == ".docx":
        from docx import Document
        from docx.text.paragraph import Paragraph
        zip_guard(filename)
        for number, block in enumerate(Document(filename).iter_inner_content(), 1):
            if isinstance(block, Paragraph):
                paragraphs(block.text, f"正文块 {number}，")
            else:
                table(([item.text for item in row.cells] for row in block.rows), f"表格 {number}，")
        notes.append("已读取 Word 正文和表格；图片与批注未作为正文提取。")
    elif extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(filename)
        if reader.is_encrypted and not reader.decrypt(""):
            raise PreviewError("请上传未加密的 PDF。")
        for number, page in enumerate(reader.pages, 1):
            paragraphs(page.extract_text() or "", f"第 {number} 页，")
        notes.append("已提取 PDF 全部页面文字；扫描图片未进行 OCR。")
    elif extension in {".json", ".jsonl"}:
        text = decode(filename)
        data = [json.loads(line) for line in text.splitlines() if line.strip()] if extension == ".jsonl" else json.loads(text)
        if isinstance(data, dict):
            candidates = [value for key, value in data.items() if normalized(key) in {"data", "questions", "items", "records", "dataset"} and isinstance(value, list)]
            if len(candidates) == 1:
                data = candidates[0]
        records = data if isinstance(data, list) else [data]
        for number, record in enumerate(records, 1):
            original_id = next((cell(value) for key, value in record.items() if normalized(key) in {"id", "题目id", "问题id", "编号", "source"} and isinstance(value, (str, int))), "") if isinstance(record, dict) else ""
            add(json.dumps(record, ensure_ascii=False), f"第 {number} 条记录", original_id)
    elif extension in TEXT_EXT:
        paragraphs(decode(filename), "")
    else:
        raise PreviewError("此格式暂不支持 AI 生成，请将旧版 Word/Excel 转为 DOCX/XLSX，或上传可读取的文字文档。")
    if not units:
        raise PreviewError("文件中没有可读取的文字或数据；扫描件请先识别成文字。")
    digest = hashlib.sha256()
    with open(filename, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return {"ok": True, "units": units, "characters": total, "sha256": digest.hexdigest(), "notes": notes}


if __name__ == "__main__":
    try:
        print(json.dumps(main(), ensure_ascii=False))
    except (PreviewError, UnicodeError) as error:
        print(json.dumps({"ok": False, "message": str(error)}, ensure_ascii=False))
    except ModuleNotFoundError:
        print(json.dumps({"ok": False, "message": "文件读取缺少 Python 依赖，请检查 requirements.txt。"}, ensure_ascii=False))
    except Exception:
        print(json.dumps({"ok": False, "message": "无法读取源文档，请检查文件格式、密码保护或损坏情况。"}, ensure_ascii=False))
