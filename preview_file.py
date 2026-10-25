"""Bounded, read-only previews. Uploaded contents are data, never instructions."""
import csv
import codecs
import datetime
import io
import json
import pathlib
import re
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8")
MAX_TEXT = 200000
MAX_ROWS = 200
MAX_COLS = 60
TEXT_EXT = {".txt", ".md", ".markdown", ".json", ".jsonl", ".log", ".yaml", ".yml", ".xml", ".html", ".htm", ".sql", ".py", ".js", ".css", ".ini", ".cfg"}


class PreviewError(Exception):
    pass


def text_content(filename):
    with open(filename, "rb") as stream:
        raw = stream.read(2 * 1024 * 1024 + 1)
    clipped = len(raw) > 2 * 1024 * 1024
    raw = raw[:2 * 1024 * 1024]
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        encoding = "utf-32"
    elif raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    elif raw.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    else:
        try:
            codecs.getincrementaldecoder("utf-8")(errors="strict").decode(raw, final=not clipped)
            encoding = "utf-8"
        except UnicodeDecodeError:
            encoding = "gb18030"
    content = raw.decode(encoding, errors="replace")
    if "\x00" in content or (content and content.count("\ufffd") > max(3, len(content) // 100)):
        raise PreviewError("文件无法按文本读取，可能是二进制文件或编码不兼容。可下载原文件查看。")
    return content, encoding, clipped


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime.datetime) else value.isoformat()
    return str(value)[:2000]


def zip_guard(filename):
    with zipfile.ZipFile(filename) as archive:
        info = archive.infolist()
        if len(info) > 15000 or sum(item.file_size for item in info) > 160 * 1024 * 1024:
            raise PreviewError("文件展开后超过预览范围，原文件可正常下载。")
        if any(item.flag_bits & 1 for item in info):
            raise PreviewError("加密文档暂不支持预览，请上传未加密副本。")
        if any(item.file_size > 80 * 1024 * 1024 for item in info):
            raise PreviewError("文档内部内容较大，超过预览范围，请下载原文件查看。")


def preview_csv(filename, extension):
    text, encoding, byte_clipped = text_content(filename)
    try:
        dialect = csv.Sniffer().sniff(text[:16000], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel_tab if extension == ".tsv" else csv.excel
    rows = []
    truncated = byte_clipped
    columns_clipped = False
    budget = 0
    csv.field_size_limit(2 * 1024 * 1024)
    for number, row in enumerate(csv.reader(io.StringIO(text, newline=""), dialect)):
        if number >= MAX_ROWS:
            truncated = True
            break
        columns_clipped |= len(row) > MAX_COLS
        cells = [safe_cell(cell) for cell in row[:MAX_COLS]]
        columns_clipped |= any(len(cell) > 2000 for cell in row[:MAX_COLS])
        budget += sum(len(cell) for cell in cells)
        if budget > MAX_TEXT:
            truncated = True
            break
        rows.append(cells)
    return {"kind": "table", "rows": rows, "rowCount": None if truncated else len(rows), "columnCount": max((len(row) for row in rows), default=0), "truncated": truncated or columns_clipped, "encoding": encoding, "sheets": [], "sheet": None}


def preview_excel(filename, sheet):
    from openpyxl import load_workbook
    zip_guard(filename)
    with open(filename, "rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            names = workbook.sheetnames
            if not names:
                raise PreviewError("工作簿中没有工作表。")
            selected = sheet if sheet is not None else names[0]
            if selected not in names:
                raise PreviewError("所选工作表不存在，请重新打开文件。")
            ws = workbook[selected]
            row_count, col_count = ws.max_row, ws.max_column
            if row_count == 1 and col_count == 1:
                ws.reset_dimensions()
                row_count = col_count = None
            rows = []
            budget = 0
            text_clipped = False
            for row in ws.iter_rows(max_row=MAX_ROWS + 1, max_col=min(col_count or MAX_COLS, MAX_COLS), values_only=True):
                text_clipped |= any(isinstance(cell, str) and len(cell) > 2000 for cell in row)
                cells = [safe_cell(cell) for cell in row]
                budget += sum(len(cell) for cell in cells)
                if budget > MAX_TEXT:
                    text_clipped = True
                    break
                rows.append(cells)
            more = len(rows) > MAX_ROWS
            rows = rows[:MAX_ROWS]
            while rows and not any(rows[-1]):
                rows.pop()
            width = max((i + 1 for row in rows for i, cell in enumerate(row) if cell), default=0)
            rows = [row[:width] for row in rows]
            return {"kind": "table", "rows": rows, "rowCount": row_count, "columnCount": col_count or width, "truncated": more or text_clipped or (row_count or 0) > MAX_ROWS or (col_count or 0) > MAX_COLS, "sheets": names[:100], "sheet": selected, "sheetCount": len(names)}
        finally:
            workbook.close()


def preview_word(filename):
    from docx import Document
    from docx.text.paragraph import Paragraph
    zip_guard(filename)
    document = Document(filename)
    blocks, budget, truncated = [], 0, False
    for item in document.iter_inner_content():
        if len(blocks) >= 500 or budget >= MAX_TEXT:
            truncated = True
            break
        if isinstance(item, Paragraph):
            text = item.text[:min(20000, MAX_TEXT - budget)]
            budget += len(text)
            if len(text) < len(item.text):
                truncated = True
            heading = re.search(r"(?:Heading|标题)\s*([1-6])", item.style.name or "", re.I)
            blocks.append({"type": "heading" if heading else "paragraph", "level": int(heading.group(1)) if heading else None, "text": text})
        else:
            rows = []
            for i, row in enumerate(item.rows):
                if i >= MAX_ROWS or budget >= MAX_TEXT:
                    truncated = True
                    break
                cells = []
                for cell in row.cells[:MAX_COLS]:
                    text = cell.text[:min(2000, max(0, MAX_TEXT - budget))]
                    budget += len(text)
                    cells.append(text)
                rows.append(cells)
            blocks.append({"type": "table", "rows": rows})
    return {"kind": "document", "blocks": blocks, "truncated": truncated, "message": "正文预览包含段落和表格，图片、批注及页面布局请下载原文件查看。"}


def preview_pdf(filename):
    from pypdf import PdfReader
    reader = PdfReader(filename)
    if reader.is_encrypted and not reader.decrypt(""):
        raise PreviewError("加密 PDF 暂不支持预览，请上传未加密副本。")
    text, truncated = [], len(reader.pages) > 20
    budget = 0
    for index, page in enumerate(reader.pages[:20]):
        content = (page.extract_text() or "")[:max(0, MAX_TEXT - budget)]
        budget += len(content)
        text.append("—— 第 %d 页 ——\n%s" % (index + 1, content))
        if budget >= MAX_TEXT:
            truncated = True
            break
    return {"kind": "text", "text": "\n\n".join(text), "pageCount": len(reader.pages), "truncated": truncated, "message": "PDF 文字预览；扫描图片可能没有可提取文字，完整排版请下载原文件查看。"}


def main():
    filename, original_name = sys.argv[1:3]
    sheet = sys.argv[3] if len(sys.argv) > 3 else None
    extension = pathlib.Path(original_name).suffix.lower()
    if extension in {".csv", ".tsv"}:
        return preview_csv(filename, extension)
    if extension in {".xlsx", ".xlsm"}:
        return preview_excel(filename, sheet)
    if extension == ".docx":
        return preview_word(filename)
    if extension == ".pdf":
        return preview_pdf(filename)
    if extension in TEXT_EXT:
        text, encoding, byte_clipped = text_content(filename)
        return {"kind": "markdown" if extension in {".md", ".markdown"} else "text", "text": text[:MAX_TEXT], "encoding": encoding, "truncated": byte_clipped or len(text) > MAX_TEXT}
    legacy = {".doc": "旧版 .doc 暂不支持正文预览，可另存为 .docx 后上传。", ".xls": "旧版 .xls 暂不支持表格预览，可另存为 .xlsx 后上传。"}
    return {"kind": "unsupported", "message": legacy.get(extension, "原文件已保存，此格式暂不支持内容预览，可下载查看。"), "truncated": False}


if __name__ == "__main__":
    try:
        print(json.dumps({"ok": True, **main()}, ensure_ascii=False))
    except ModuleNotFoundError:
        print(json.dumps({"ok": False, "message": "当前 Python 环境缺少此格式的预览依赖，请检查工具的 Python 配置；原文件可下载。"}, ensure_ascii=False))
    except PreviewError as exc:
        print(json.dumps({"ok": False, "message": str(exc)}, ensure_ascii=False))
    except Exception:
        print(json.dumps({"ok": False, "message": "无法解析此文件，可能已损坏、受密码保护或格式与扩展名不一致；原文件可下载。"}, ensure_ascii=False))
